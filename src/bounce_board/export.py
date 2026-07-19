"""
Bounce Board — report export (markdown / PDF / DOCX).

Builds the deliverable from the structured report on the session row rather
than converting rendered HTML, so exports work headlessly and stay consistent
with the in-app report. PDF via fpdf2 (pure python), DOCX via python-docx.
"""
import io
import re
from datetime import datetime

_MD_BOLD = re.compile(r"\*\*(.+?)\*\*")
_MD_HEADING = re.compile(r"^#{1,6}\s*", re.MULTILINE)


def _plain(md: str, drop_title: str = "") -> str:
    """Markdown → plain text good enough for PDF/DOCX paragraphs."""
    text = _MD_HEADING.sub("", md or "")
    text = _MD_BOLD.sub(r"\1", text)
    text = text.replace("*", "").strip()
    # The section markdown usually starts with its own title — we print our
    # own heading, so drop the duplicate first line.
    if drop_title and text.lower().startswith(drop_title.lower()):
        text = text[len(drop_title):].lstrip()
    return text


def _money(n: float) -> str:
    if n >= 1_000_000:
        return f"${n / 1_000_000:.1f}M"
    if n >= 1_000:
        return f"${round(n / 1_000)}k"
    return f"${n:.0f}"


def _report_sections(row: dict) -> dict:
    report = row.get("report") or {}
    return {
        "title": row.get("title") or "Bounce Board analysis",
        "generated_at": report.get("generatedAt") or "",
        "executive_summary": report.get("executiveSummaryMd") or "",
        "recommendations": report.get("recommendations") or row.get("recommendations") or [],
        "action_plan": report.get("actionPlan") or [],
        "cost_roi": report.get("costRoi") or {},
        "risk_register": report.get("riskRegister") or [],
        "management_report": report.get("managementReportMd") or "",
    }


def _generated_line(sections: dict) -> str:
    try:
        stamp = datetime.fromisoformat(sections["generated_at"]).strftime("%Y-%m-%d %H:%M")
    except (TypeError, ValueError):
        stamp = sections["generated_at"]
    return f"Bounce Board AI executive analysis - generated {stamp}"


def filename(row: dict, ext: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", (row.get("title") or "report").lower()).strip("-")[:60]
    return f"{slug or 'report'}-report.{ext}"


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def to_markdown(row: dict) -> str:
    s = _report_sections(row)
    cost = s["cost_roi"]
    lines: list[str] = [
        f"# {s['title']}",
        "",
        f"*{_generated_line(s)}*",
        "",
        s["executive_summary"],
        "",
        "## Recommendations",
        "",
        "| # | Recommendation | Composite | Cost | ROI | Effort | Red team |",
        "|---|---|---|---|---|---|---|",
    ]
    for i, r in enumerate(s["recommendations"]):
        verdict = (r.get("critique") or {}).get("verdict", "-")
        lines.append(
            f"| {i + 1} | {r.get('title')} | {r.get('compositeScore')} | "
            f"{_money(r.get('estimatedCost') or 0)} | {r.get('estimatedRoiPct')}% | "
            f"{r.get('effortWeeks')}w | {verdict} |"
        )
    lines += ["", "## Action plan", ""]
    for a in s["action_plan"]:
        end = (a.get("startWeek") or 1) + (a.get("durationWeeks") or 1) - 1
        lines.append(
            f"- **{a.get('title')}** — {a.get('ownerRole')}, weeks {a.get('startWeek')}–{end} ({a.get('phase')})"
        )
    lines += [
        "",
        "## Cost & ROI (estimates)",
        "",
        f"Total investment {_money(cost.get('totalCost') or 0)}, expected ROI "
        f"{cost.get('expectedRoiPct')}%, payback {cost.get('paybackMonths')} months.",
        "",
    ]
    if cost.get("assumptions"):
        lines.append("Assumptions:")
        lines += [f"- {a}" for a in cost["assumptions"]]
        lines.append("")
    lines += ["## Risk register", ""]
    for r in s["risk_register"]:
        lines.append(
            f"- **{r.get('title')}** ({r.get('severity')}) — owner {r.get('owner')}. {r.get('mitigation')}"
        )
    lines += ["", s["management_report"]]
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# PDF (fpdf2)
# ---------------------------------------------------------------------------

_LATIN_FIXES = {
    "–": "-", "—": "-", "‘": "'", "’": "'",
    "“": '"', "”": '"', "…": "...", "→": "->", "·": "-",
}


def _latin(text: str) -> str:
    for src, dst in _LATIN_FIXES.items():
        text = text.replace(src, dst)
    return text.encode("latin-1", "replace").decode("latin-1")


def to_pdf(row: dict) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    s = _report_sections(row)
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    # multi_cell defaults leave the cursor at the cell's right edge; every
    # paragraph here is full-width, so always return to the left margin.
    flow = {"new_x": XPos.LMARGIN, "new_y": YPos.NEXT}

    def heading(text: str, size: int = 13, top: int = 6) -> None:
        pdf.ln(top)
        pdf.set_font("helvetica", "B", size)
        pdf.multi_cell(usable, 7, _latin(text), **flow)
        pdf.ln(1)

    def body(text: str) -> None:
        pdf.set_font("helvetica", "", 10)
        pdf.multi_cell(usable, 5.5, _latin(text), **flow)

    def table(headers: list[str], widths: list[float], rows: list[list[str]]) -> None:
        pdf.set_font("helvetica", "B", 8.5)
        pdf.set_fill_color(240, 240, 240)
        for h, w in zip(headers, widths):
            pdf.cell(w, 6, _latin(h), border=1, fill=True)
        pdf.ln()
        pdf.set_font("helvetica", "", 8.5)
        for cells in rows:
            # First column may wrap — measure its height and draw the row box.
            first = _latin(str(cells[0]))
            lines = pdf.multi_cell(widths[0], 5, first, dry_run=True, output="LINES")
            height = max(5 * len(lines), 5)
            if pdf.get_y() + height > pdf.page_break_trigger:
                pdf.add_page()
            x, y = pdf.get_x(), pdf.get_y()
            pdf.multi_cell(widths[0], 5, first, border=1, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.set_xy(x + widths[0], y)
            for value, w in zip(cells[1:], widths[1:]):
                pdf.cell(w, height, _latin(str(value)), border=1)
            pdf.ln(height)

    # Header
    pdf.set_font("helvetica", "B", 17)
    pdf.multi_cell(usable, 9, _latin(s["title"]), **flow)
    pdf.set_font("helvetica", "I", 9)
    pdf.set_text_color(110, 110, 110)
    pdf.multi_cell(usable, 5, _latin(_generated_line(s)), **flow)
    pdf.set_text_color(0, 0, 0)

    heading("Executive Summary")
    body(_plain(s["executive_summary"], drop_title="Executive Summary"))

    heading("Ranked Recommendations")
    table(
        ["Recommendation", "Score", "Cost", "ROI", "Effort", "Red team"],
        [usable - 82, 15, 20, 15, 15, 17],
        [
            [
                r.get("title") or "",
                str(r.get("compositeScore") or ""),
                _money(r.get("estimatedCost") or 0),
                f"{r.get('estimatedRoiPct') or 0}%",
                f"{r.get('effortWeeks') or 0}w",
                (r.get("critique") or {}).get("verdict", "-"),
            ]
            for r in s["recommendations"]
        ],
    )

    if s["action_plan"]:
        heading("Action Plan")
        for a in s["action_plan"]:
            end = (a.get("startWeek") or 1) + (a.get("durationWeeks") or 1) - 1
            body(f"- {a.get('title')} — {a.get('ownerRole')}, weeks {a.get('startWeek')}-{end} ({a.get('phase')})")

    cost = s["cost_roi"]
    heading("Cost & ROI (estimates)")
    body(
        f"Total investment {_money(cost.get('totalCost') or 0)} - expected ROI "
        f"{cost.get('expectedRoiPct') or 0}% - payback {cost.get('paybackMonths') or 0} months"
    )
    for b in cost.get("breakdown") or []:
        body(f"- {b.get('label')}: {_money(b.get('amount') or 0)}")
    if cost.get("assumptions"):
        pdf.set_font("helvetica", "I", 9)
        pdf.multi_cell(usable, 5, _latin("Assumptions: " + "; ".join(cost["assumptions"])), **flow)

    if s["risk_register"]:
        heading("Risk Register")
        table(
            ["Risk", "Severity", "Owner", "Status"],
            [usable - 75, 22, 33, 20],
            [
                [r.get("title") or "", r.get("severity") or "", r.get("owner") or "", r.get("status") or ""]
                for r in s["risk_register"]
            ],
        )

    heading("Management Report")
    body(_plain(s["management_report"], drop_title="Management Report"))

    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------


def to_docx(row: dict) -> bytes:
    from docx import Document

    s = _report_sections(row)
    doc = Document()
    doc.add_heading(s["title"], level=0)
    doc.add_paragraph(_generated_line(s)).italic = True

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(_plain(s["executive_summary"], drop_title="Executive Summary"))

    doc.add_heading("Ranked Recommendations", level=1)
    headers = ["#", "Recommendation", "Score", "Cost", "ROI", "Effort", "Red team"]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for cell, h in zip(t.rows[0].cells, headers):
        cell.text = h
    for i, r in enumerate(s["recommendations"]):
        cells = t.add_row().cells
        values = [
            str(i + 1),
            r.get("title") or "",
            str(r.get("compositeScore") or ""),
            _money(r.get("estimatedCost") or 0),
            f"{r.get('estimatedRoiPct') or 0}%",
            f"{r.get('effortWeeks') or 0}w",
            (r.get("critique") or {}).get("verdict", "-"),
        ]
        for cell, v in zip(cells, values):
            cell.text = v

    if s["action_plan"]:
        doc.add_heading("Action Plan", level=1)
        for a in s["action_plan"]:
            end = (a.get("startWeek") or 1) + (a.get("durationWeeks") or 1) - 1
            doc.add_paragraph(
                f"{a.get('title')} — {a.get('ownerRole')}, weeks {a.get('startWeek')}–{end} ({a.get('phase')})",
                style="List Bullet",
            )

    cost = s["cost_roi"]
    doc.add_heading("Cost & ROI (estimates)", level=1)
    doc.add_paragraph(
        f"Total investment {_money(cost.get('totalCost') or 0)} · expected ROI "
        f"{cost.get('expectedRoiPct') or 0}% · payback {cost.get('paybackMonths') or 0} months"
    )
    for b in cost.get("breakdown") or []:
        doc.add_paragraph(f"{b.get('label')}: {_money(b.get('amount') or 0)}", style="List Bullet")
    if cost.get("assumptions"):
        doc.add_paragraph("Assumptions: " + "; ".join(cost["assumptions"])).italic = True

    if s["risk_register"]:
        doc.add_heading("Risk Register", level=1)
        rt = doc.add_table(rows=1, cols=4)
        rt.style = "Light Grid Accent 1"
        for cell, h in zip(rt.rows[0].cells, ["Risk", "Severity", "Owner", "Mitigation"]):
            cell.text = h
        for r in s["risk_register"]:
            cells = rt.add_row().cells
            for cell, v in zip(
                cells,
                [r.get("title") or "", r.get("severity") or "", r.get("owner") or "", r.get("mitigation") or ""],
            ):
                cell.text = v

    doc.add_heading("Management Report", level=1)
    doc.add_paragraph(_plain(s["management_report"], drop_title="Management Report"))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
